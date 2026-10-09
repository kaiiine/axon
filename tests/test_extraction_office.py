"""`.pptx`, `.xlsx`, `.docx` sont des fichiers ZIP — jamais du texte UTF-8.

Vécu : un `.pptx` de 23 Mo lu comme du texte a donné 100 000 caractères de
« � » à l'agent, deux fois (l'attache puis `local_read_file`), avant qu'il
improvise `unzip -l` en shell — bloqué par la confirmation. `.pdf` avait déjà
son extracteur dédié (`_extract_pdf`) ; les trois autres n'avaient rien.

Les fichiers ici sont générés à la volée avec `python-pptx`, `openpyxl` et
`python-docx` — le contenu attendu est donc connu avec certitude, pas supposé.
"""
from __future__ import annotations

import pytest

pptx = pytest.importorskip("pptx")
openpyxl = pytest.importorskip("openpyxl")
docx = pytest.importorskip("docx")


@pytest.fixture
def un_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    diapo1 = prs.slides.add_slide(prs.slide_layouts[1])
    diapo1.shapes.title.text = "Critères d'audit RGAA"
    diapo1.placeholders[1].text = "19 critères à vérifier"

    diapo2 = prs.slides.add_slide(prs.slide_layouts[6])
    tbl = diapo2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "Critère"
    tbl.cell(0, 1).text = "Test"
    tbl.cell(1, 0).text = "1.1 Images"
    tbl.cell(1, 1).text = "Alt text présent"

    chemin = tmp_path / "audit.pptx"
    prs.save(str(chemin))
    return chemin


@pytest.fixture
def un_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Grille RGAA"
    ws.append(["Critère", "Sous-critère"])
    ws.append(["1", "1.1 Images informatives"])

    chemin = tmp_path / "grille.xlsx"
    wb.save(str(chemin))
    return chemin


@pytest.fixture
def un_docx(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("Rapport d'audit RGAA", level=1)
    doc.add_paragraph("Contexte de l'audit accessibilité.")

    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Critère"
    t.cell(0, 1).text = "Statut"
    t.cell(1, 0).text = "1.1 Images"
    t.cell(1, 1).text = "Conforme"

    doc.add_paragraph("Suite du texte après le tableau.")

    chemin = tmp_path / "audit.docx"
    doc.save(str(chemin))
    return chemin


# ── extraction directe ─────────────────────────────────────────────────────
def test_le_texte_des_diapositives_est_extrait(un_pptx):
    from src.ui.attachments import _extract_pptx

    texte = _extract_pptx(un_pptx)

    assert "Critères d'audit RGAA" in texte
    assert "19 critères à vérifier" in texte
    assert "[Diapositive 1]" in texte and "[Diapositive 2]" in texte


def test_le_tableau_dune_diapositive_est_extrait(un_pptx):
    from src.ui.attachments import _extract_pptx

    texte = _extract_pptx(un_pptx)

    assert "Critère | Test" in texte
    assert "1.1 Images | Alt text présent" in texte


def test_les_cellules_dune_feuille_sont_extraites(un_xlsx):
    from src.ui.attachments import _extract_xlsx

    texte = _extract_xlsx(un_xlsx)

    assert "[Feuille : Grille RGAA]" in texte
    assert "1 | 1.1 Images informatives" in texte


def test_les_titres_et_paragraphes_dun_docx_sont_extraits(un_docx):
    from src.ui.attachments import _extract_docx

    texte = _extract_docx(un_docx)

    assert "## Rapport d'audit RGAA" in texte
    assert "Contexte de l'audit accessibilité." in texte


def test_le_tableau_dun_docx_reste_a_sa_place(un_docx):
    """Le point qui compte : `python-docx` expose paragraphes et tableaux comme
    deux collections séparées. Les recoller dans l'ordre du XML, et pas
    paragraphes-puis-tableaux, évite qu'un audit qui dit « le tableau ci-dessus »
    perde son référent."""
    from src.ui.attachments import _extract_docx

    lignes = _extract_docx(un_docx).splitlines()
    i_contexte = next(i for i, l in enumerate(lignes) if "Contexte" in l)
    i_tableau = next(i for i, l in enumerate(lignes) if "Critère | Statut" in l)
    i_suite = next(i for i, l in enumerate(lignes) if "Suite du texte" in l)

    assert i_contexte < i_tableau < i_suite


def test_un_pptx_illisible_ne_leve_pas(tmp_path):
    from src.ui.attachments import _extract_pptx

    faux = tmp_path / "casse.pptx"
    faux.write_bytes(b"pas un vrai zip")

    assert "Erreur" in _extract_pptx(faux)


def test_un_docx_illisible_ne_leve_pas(tmp_path):
    from src.ui.attachments import _extract_docx

    faux = tmp_path / "casse.docx"
    faux.write_bytes(b"pas un vrai zip")

    assert "Erreur" in _extract_docx(faux)


# ── ce que l'ancien code produisait, comparé ────────────────────────────────
def test_lire_en_texte_brut_donnait_du_bruit(un_pptx):
    """La preuve du bug, pas juste son récit. Un ZIP contient assez d'octets
    ASCII (ses en-têtes, ses noms de fichiers XML internes) pour que le taux
    de « � » seul ne soit pas un critère fiable sur un petit fichier — la
    preuve qui compte est que le VRAI contenu, lui, n'y est nulle part."""
    brut = un_pptx.read_text(encoding="utf-8", errors="replace")

    assert "Critères d'audit RGAA" not in brut
    assert "�" in brut


# ── les deux points d'entrée utilisent le même extracteur ──────────────────
def test_add_text_appelle_lextracteur_pptx(un_pptx):
    from src.ui.attachments import AttachmentStore

    store = AttachmentStore()
    piece = store._add_text(un_pptx)

    assert "Critères d'audit RGAA" in piece.content
    assert "�" not in piece.content


def test_local_read_file_appelle_le_meme_extracteur(un_pptx):
    from src.agents.filesystem.tools import local_read_file

    resultat = local_read_file.invoke({"path": str(un_pptx)})

    assert resultat["status"] == "ok"
    assert "Critères d'audit RGAA" in resultat["content"]
    assert "�" not in resultat["content"]


def test_local_read_file_sur_xlsx(un_xlsx):
    from src.agents.filesystem.tools import local_read_file

    resultat = local_read_file.invoke({"path": str(un_xlsx)})

    assert resultat["status"] == "ok"
    assert "1.1 Images informatives" in resultat["content"]


def test_local_read_file_sur_docx(un_docx):
    from src.agents.filesystem.tools import local_read_file

    resultat = local_read_file.invoke({"path": str(un_docx)})

    assert resultat["status"] == "ok"
    assert "Rapport d'audit RGAA" in resultat["content"]


def test_les_extensions_couvertes_sont_toutes_enregistrees():
    from src.ui.attachments import EXTRACTEURS_OFFICE

    assert {".pptx", ".xlsx", ".docx"} <= set(EXTRACTEURS_OFFICE)
