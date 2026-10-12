from __future__ import annotations

import base64
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, List, Optional

MAX_TEXT_SIZE         = 100_000  # extraction brute complète
_ORCHESTRATOR_INJECT  =  20_000  # seuil au-delà duquel on n'injecte pas le contenu dans la conversation

_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
_MIME = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
         "gif": "image/gif", "webp": "image/webp"}


def _fix_spaced_chars(text: str) -> str:
    """Corrige l'artefact pypdf où les lettres sont séparées par des espaces.
    """
    import re
    lines = []
    for line in text.split('\n'):
        words = line.split(' ')
        single = sum(1 for w in words if len(w) <= 1)
        # Si >60% des "mots" sont des lettres seules → ligne espacée
        if len(words) > 4 and single / len(words) > 0.6:
            line = re.sub(r'(?<=[A-Za-zÀ-ÿ0-9]) (?=[A-Za-zÀ-ÿ0-9])', '', line)
        lines.append(line)
    return '\n'.join(lines)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = _fix_spaced_chars(text)
            if text.strip():
                pages.append(f"[Page {i + 1}]\n{text}")
        return "\n\n".join(pages)[:MAX_TEXT_SIZE] or "[PDF sans texte extractible]"
    except Exception as e:
        return f"[Erreur lecture PDF: {e}]"


def _extract_pptx(path: Path) -> str:
    """Le texte d'un .pptx — un fichier ZIP, jamais du texte UTF-8.

    Vécu : un `.pptx` de 23 Mo lu comme du texte donnait 100 000 caractères de
    `�` — un fichier binaire qui n'a que « texte » comme catégorie ici, faute
    d'un extracteur dédié comme pour le `.pdf`. Le modèle a improvisé un
    `unzip -l` en shell, bloqué par la confirmation.
    """
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        diapos = []
        for i, diapo in enumerate(prs.slides, 1):
            morceaux = []
            for forme in diapo.shapes:
                if forme.has_text_frame and forme.text_frame.text.strip():
                    morceaux.append(forme.text_frame.text.strip())
                if forme.has_table:
                    for ligne in forme.table.rows:
                        morceaux.append(" | ".join(c.text.strip() for c in ligne.cells))
            notes = diapo.notes_slide.notes_text_frame.text.strip() \
                if diapo.has_notes_slide else ""
            bloc = "\n".join(morceaux) or "(diapositive sans texte)"
            if notes:
                bloc += f"\n[Notes] {notes}"
            diapos.append(f"[Diapositive {i}]\n{bloc}")
        return "\n\n".join(diapos)[:MAX_TEXT_SIZE] or "[.pptx sans texte extractible]"
    except Exception as e:
        return f"[Erreur lecture .pptx: {e}]"


def _extract_xlsx(path: Path) -> str:
    """Le texte d'un .xlsx — même famille que le .pptx, même défaut sans lui.

    Une feuille immense est TRONQUÉE par lignes, jamais par colonnes : perdre
    la fin d'un tableau laisse au moins des lignes complètes, perdre des
    colonnes casserait chaque ligne restante.
    """
    try:
        from openpyxl import load_workbook
        classeur = load_workbook(str(path), read_only=True, data_only=True)
        feuilles = []
        for nom in classeur.sheetnames:
            feuille = classeur[nom]
            lignes = []
            for ligne in feuille.iter_rows(values_only=True):
                if any(c is not None for c in ligne):
                    lignes.append(" | ".join(
                        "" if c is None else str(c) for c in ligne))
                if len(lignes) >= 500:
                    lignes.append(f"[… tronqué, {feuille.max_row} lignes au total]")
                    break
            feuilles.append(f"[Feuille : {nom}]\n" + "\n".join(lignes))
        return "\n\n".join(feuilles)[:MAX_TEXT_SIZE] or "[.xlsx sans contenu extractible]"
    except Exception as e:
        return f"[Erreur lecture .xlsx: {e}]"


def _extract_docx(path: Path) -> str:
    """Le texte d'un .docx — même famille, même défaut sans extracteur dédié.

    Paragraphes et tableaux sont rendus DANS L'ORDRE du document : `python-docx`
    les expose comme deux collections séparées (`.paragraphs`, `.tables`), et
    les recoller dans l'ordre du XML (`document.element.body`) est ce qui évite
    qu'un tableau atterrisse après tout le texte qui le suit vraiment — un
    audit qui commente « le tableau ci-dessus » perdrait son référent.
    """
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(str(path))
        morceaux = []
        for enfant in doc.element.body.iterchildren():
            if enfant.tag.endswith("}p"):
                p = Paragraph(enfant, doc)
                if not p.text.strip():
                    continue
                prefixe = "## " if p.style and p.style.name.startswith("Heading") else ""
                morceaux.append(f"{prefixe}{p.text.strip()}")
            elif enfant.tag.endswith("}tbl"):
                t = Table(enfant, doc)
                for ligne in t.rows:
                    morceaux.append(" | ".join(c.text.strip() for c in ligne.cells))
        return "\n".join(morceaux)[:MAX_TEXT_SIZE] or "[.docx sans texte extractible]"
    except Exception as e:
        return f"[Erreur lecture .docx: {e}]"


#: Un extracteur par extension, appelé par `_add_text` ET par `local_read_file` —
#: un seul endroit qui sait comment lire chaque format, pas deux qui divergent.
EXTRACTEURS_OFFICE: dict[str, Callable[[Path], str]] = {
    ".pptx": _extract_pptx,
    ".ppsx": _extract_pptx,
    ".xlsx": _extract_xlsx,
    ".xlsm": _extract_xlsx,
    ".docx": _extract_docx,
}


@dataclass
class Attachment:
    name: str
    is_image: bool
    content: str = ""        # fichiers texte (extrait complet)
    b64: str = ""            # images
    mime: str = ""
    size_hint: str = ""
    source_path: str = ""    # chemin absolu d'origine (pour re-lecture par les tools)


class AttachmentStore:
    def __init__(self):
        self._items: List[Attachment] = []
        self._derniers: List[Attachment] = []

    def add_file(self, path_str: str) -> Optional[Attachment]:
        path = Path(path_str)
        if not path.exists():
            return None
        if path.suffix.lower() in _IMAGE_EXTS:
            return self._add_image(path)
        return self._add_text(path)

    def add_clipboard_image(self, img) -> Attachment:
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        b64 = base64.b64encode(buf.getvalue()).decode()
        w, h = img.size
        a = Attachment(name="clipboard.png", is_image=True,
                       b64=b64, mime="image/png", size_hint=f"{w}×{h}")
        self._items.append(a)
        return a

    def _add_image(self, path: Path) -> Attachment:
        data = path.read_bytes()
        b64 = base64.b64encode(data).decode()
        mime = _MIME.get(path.suffix.lstrip(".").lower(), "image/png")
        size_hint = f"{len(data) // 1024}KB"
        try:
            from PIL import Image
            with Image.open(path) as img:
                size_hint = f"{img.width}×{img.height}"
        except Exception:
            pass
        a = Attachment(name=path.name, is_image=True, b64=b64,
                       mime=mime, size_hint=size_hint)
        self._items.append(a)
        return a

    def _add_text(self, path: Path) -> Attachment:
        size = path.stat().st_size
        size_hint = f"{size // 1024}KB" if size >= 1024 else f"{size}B"
        suffixe = path.suffix.lower()
        if suffixe == ".pdf":
            content = _extract_pdf(path)
        elif suffixe in EXTRACTEURS_OFFICE:
            content = EXTRACTEURS_OFFICE[suffixe](path)
        else:
            try:
                content = path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT_SIZE]
            except Exception as e:
                content = f"[erreur: {e}]"
        a = Attachment(name=path.name, is_image=False,
                       content=content, size_hint=size_hint,
                       source_path=str(path.resolve()))
        self._items.append(a)
        return a

    def pop_all(self) -> List[Attachment]:
        items = list(self._items)
        self._items.clear()
        # Retenues pour le tour : l'agent de code s'exécute APRÈS que l'UI ait
        # vidé la pile pour bâtir le message de l'orchestrateur, et il ne reçoit
        # qu'une chaîne de tâche. Sans ça, joindre un PDF puis demander « code ce
        # qui est décrit dedans » n'envoyait au sous-graphe que la phrase.
        if items:
            self._derniers = items
        return items

    @property
    def derniers(self) -> List[Attachment]:
        """Les pièces jointes du dernier tour, texte compris."""
        return list(self._derniers)

    def remove(self, name: str) -> bool:
        before = len(self._items)
        self._items = [a for a in self._items if a.name != name]
        return len(self._items) < before

    @property
    def items(self) -> List[Attachment]:
        return list(self._items)

    def __len__(self):
        return len(self._items)

    def __bool__(self):
        return bool(self._items)


def get_clipboard_image():
    """Retourne une PIL Image depuis le presse-papiers (Wayland + X11)."""
    import subprocess, io

    # Wayland : wl-paste
    try:
        r = subprocess.run(["wl-paste", "--type", "image/png"],
                           capture_output=True, timeout=3)
        if r.returncode == 0 and r.stdout:
            from PIL import Image
            return Image.open(io.BytesIO(r.stdout))
    except (FileNotFoundError, subprocess.TimeoutExpired, Exception):
        pass

    # X11 : xclip
    try:
        r = subprocess.run(["xclip", "-selection", "clipboard", "-t", "image/png", "-o"],
                           capture_output=True, timeout=3)
        if r.returncode == 0 and r.stdout:
            from PIL import Image
            return Image.open(io.BytesIO(r.stdout))
    except (FileNotFoundError, subprocess.TimeoutExpired, Exception):
        pass

    # Fallback PIL (X11)
    try:
        from PIL import ImageGrab
        img = ImageGrab.grabclipboard()
        if img is not None:
            return img
    except Exception:
        pass

    return None


def open_file_picker() -> Optional[str]:
    """Ouvre un sélecteur de fichier natif. Retourne le chemin ou None.

    L'île (couche layer-shell TOP) passerait au-dessus du dialogue : elle est
    priée de s'effacer le temps du choix, puis de revenir. Sans île lancée,
    `pont.envoyer` échoue vite et en silence."""
    from .pont_overlay import pont

    pont.envoyer(type="selecteur_fichier", ouvert=True)
    try:
        return _ouvrir_selecteur_natif()
    finally:
        pont.envoyer(type="selecteur_fichier", ouvert=False)


def _ouvrir_selecteur_natif() -> Optional[str]:
    import subprocess

    # zenity (GNOME / Wayland)
    try:
        r = subprocess.run(
            ["zenity", "--file-selection", "--title=Joindre un fichier"],
            capture_output=True, text=True, timeout=120,
        )
        if r.returncode == 0:
            return r.stdout.strip() or None
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    # kdialog (KDE)
    try:
        r = subprocess.run(
            ["kdialog", "--getopenfilename", ".", "--title", "Joindre un fichier"],
            capture_output=True, text=True, timeout=120,
        )
        if r.returncode == 0:
            return r.stdout.strip() or None
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    # tkinter fallback
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.wm_attributes("-topmost", 1)
        root.update()
        path = filedialog.askopenfilename(title="Joindre un fichier")
        root.destroy()
        return path or None
    except Exception:
        pass

    return None


def build_message_with_attachments(text: str, attachments: List[Attachment]) -> dict:
    """
    Construit le message LangChain avec les pièces jointes injectées.
    - Fichiers texte → injectés dans le contenu textuel
    - Images → message multimodal (content list)
    """
    if not attachments:
        return {"role": "user", "content": text}

    text_parts = [text]
    images = []

    for a in attachments:
        if a.is_image:
            images.append(a)
        else:
            lang = Path(a.name).suffix.lstrip(".")
            if len(a.content) > _ORCHESTRATOR_INJECT:
                # File too large to inject: give the LLM the local path
                text_parts.append(
                    f"\n\n---\n**Fichier joint : {a.name}** — LOCAL file on disk (do NOT search on Google Drive)\n"
                    f"Absolute path: `{a.source_path}` · Size: {a.size_hint}\n"
                    f"→ Use local_read_file(path=\"{a.source_path}\") to read the content in sections."
                )
            else:
                text_parts.append(
                    f"\n\n---\n**Fichier joint : {a.name}**\n```{lang}\n{a.content}\n```"
                )

    full_text = "".join(text_parts)

    if not images:
        return {"role": "user", "content": full_text}

    content: list = [{"type": "text", "text": full_text}]
    for img in images:
        content.append({
            "type": "image_url",
            "image_url": {"url": f"data:{img.mime};base64,{img.b64}"},
        })
    return {"role": "user", "content": content}


#: La pile du tour. Elle vivait dans `streaming`, qui est le seul à la remplir —
#: mais plus le seul à la lire : le nœud `coder` y prend les pièces à transmettre
#: au sous-graphe, et importer un nom privé de `streaming` depuis l'orchestrateur
#: aurait bouclé.
attachments = AttachmentStore()
